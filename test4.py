import openpyxl as xl
import docx
import tkinter as tk
from tkinter import filedialog

janela = tk.Tk()
janela.geometry("500x400")


def selecionar_pasta():
    caminho_arquivo = filedialog.askdirectory(title="Selecione uma pasta para Salvar!")
    return caminho_arquivo

def extrair_matriculas():
    workbook = xl.load_workbook("doc.xlsx")
    planilha = workbook.active

    local = selecionar_pasta()

    for rows in planilha.iter_rows(min_row=2):
        mat = rows[0].value
        texto = rows[1].value

        doc = docx.Document()
        doc.add_paragraph(f"Matrícula: {texto}")

        nome_arquivo = f"{mat} FALTA CONFERIR.docx"
        caminho_completo = f"{local}/{nome_arquivo}"
        doc.save(caminho_completo)

    workbook.close()


texto = tk.Button(janela, text="Onde Salvar", command=extrair_matriculas)
texto.pack(padx=10, pady=10)



janela.mainloop()

