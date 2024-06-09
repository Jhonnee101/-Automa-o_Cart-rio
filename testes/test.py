import openpyxl as xl
import docx


workbook = xl.load_workbook("doc.xlsx")
planilha = workbook.active

for rows in planilha.iter_rows(min_row=2):
    mat = rows[0].value
    texto = rows[1].value


    doc = docx.Document()
    doc.add_paragraph(f"Matrícula: {texto}")

    nome_arquivo = f"{mat} FALTA CONFERIR.docx"
    doc.save(nome_arquivo)


workbook.close()


