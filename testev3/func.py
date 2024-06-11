from tkinter import messagebox
from tkinter import filedialog
import openpyxl as xl
import docx
from docx.shared import Pt

caminho_entrada = ""
caminho_saida = ""

def selecionar_entrada():
    global caminho_entrada
    arquivo = filedialog.askopenfile(title="Selecione a planilha para analisar!")
    if arquivo:
        caminho_entrada = arquivo.name

def selecionar_pasta():
    global caminho_saida
    caminho_pasta = filedialog.askdirectory(title="Selecione uma pasta para Salvar!")
    caminho_saida = caminho_pasta

def extrair_matriculas():
    workbook = xl.load_workbook(caminho_entrada)
    planilha = workbook.active

    doc = docx.Document()
    doc.styles['Normal'].font.name = 'Courier New'
    doc.styles['Normal'].font.size = Pt(12)
    doc.styles['Normal'].paragraph_format.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.JUSTIFY

    for rows in planilha.iter_rows(min_row=2):
        mat = rows[0].value
        texto = rows[1].value

        doc.add_paragraph(f"Matrícula: {texto}")

        nome_arquivo = f"{mat} FALTA CONFERIR.docx"
        caminho_completo = f"{caminho_saida}/{nome_arquivo}"
        doc.save(caminho_completo)

    workbook.close()
    messagebox.showinfo("Concluído", "Processo concluído com sucesso!")